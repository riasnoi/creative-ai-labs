#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
ROOT = REPO_ROOT / "ЛР01"
RESULTS = json.loads((ROOT / "artifacts/v13/results.json").read_text(encoding="utf-8"))
RESULTS_50 = json.loads((ROOT / "artifacts/v13_n50/results.json").read_text(encoding="utf-8"))
parser = argparse.ArgumentParser(description="Собрать черновик отчёта из сохранённых результатов варианта 13.")
parser.add_argument("--out", type=Path, default=REPO_ROOT / ".runs/report_v13.docx")
args = parser.parse_args()
OUT = args.out.resolve()
if OUT == (ROOT / "reports/Рясной_Вариант_13_Лаб1.docx").resolve():
    parser.error("сдаваемый отчёт сохранён после ручной редакции; выберите другое имя черновика")
BASE_IMAGE = ROOT / "artifacts/v13/artifacts/base_seed101.png"
FACTOR_IMAGE = ROOT / "artifacts/v13/artifacts/perturbed_seed101.png"
TITLE_LOGO = ROOT / "reports/assets/dgtu-logo.png"


def set_run_font(run, name="Times New Roman", size=11, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, left=140, bottom=100, right=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    # Use the classic OOXML left/right margin names for consistent rendering in
    # both Microsoft Word and LibreOffice. Values are twips: 100 = 5 pt,
    # 140 = 7 pt.
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7B7B7", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "none")
        node.set(qn("w:sz"), "0")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")


def style_table(table, header=True, widths=None, font_size=9, alignments=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for column, width in zip(table.columns, widths):
            column.width = width
    set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        if widths:
            for c_idx, width in enumerate(widths):
                row.cells[c_idx].width = width
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if header and r_idx == 0:
                set_cell_shading(cell, "EDEDED")
            else:
                set_cell_shading(cell, "FFFFFF")
            for p in cell.paragraphs:
                if alignments and c_idx < len(alignments):
                    p.alignment = alignments[c_idx]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=(header and r_idx == 0), color="000000")
    if header:
        set_repeat_table_header(table.rows[0])


def add_table(doc, headers, rows, widths=None, font_size=9, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = str(value)
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            row.cells[i].text = str(value)
    style_table(table, widths=widths, font_size=font_size, alignments=alignments)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_command(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, name="Courier New", size=9)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color="000000")


def add_title_line(doc, text, size=14, space_before=0, space_after=0, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.79)
section.bottom_margin = Inches(0.79)
section.left_margin = Inches(1.18)
section.right_margin = Inches(0.59)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
for list_style_name in ("List Bullet", "List Number"):
    list_style = styles[list_style_name]
    list_style.font.name = "Times New Roman"
    list_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    list_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    list_style.font.size = Pt(11)
    list_style.font.color.rgb = RGBColor(0, 0, 0)
for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
    style = styles[name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    if name == "Title":
        p_pr = style._element.get_or_add_pPr()
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)

core = doc.core_properties
core.title = "Лабораторная работа 1 Подготовка рабочей среды и паспорт эксперимента"
core.subject = "Вариант 13"
core.author = "Владимир Рясной"
core.keywords = "воспроизводимость, seed, SHA-256, паспорт эксперимента, Engee"

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Inches(0)
p.paragraph_format.space_after = Pt(4)
logo = p.add_run().add_picture(str(TITLE_LOGO), width=Inches(0.68))
logo._inline.docPr.set("descr", "Эмблема Донского государственного технического университета")
logo._inline.docPr.set("title", "Эмблема ДГТУ")

p = add_title_line(
    doc,
    "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ",
    size=12,
    space_after=12,
)
p.paragraph_format.left_indent = Inches(-0.1)
add_title_line(
    doc,
    "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ БЮДЖЕТНОЕ\n"
    "ОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ\n"
    "ВЫСШЕГО ОБРАЗОВАНИЯ\n"
    "«ДОНСКОЙ ГОСУДАРСТВЕННЫЙ ТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ»\n"
    "(ДГТУ)",
    size=14,
)

add_title_line(doc, "ЛАБОРАТОРНАЯ РАБОТА №1", size=14, space_before=96, space_after=6)
add_title_line(doc, "по дисциплине:", size=14, space_after=3)
add_title_line(doc, "«Искусственный интеллект в креативных технологиях»", size=14, space_after=22)

title_details = doc.add_table(rows=6, cols=3)
title_details.alignment = WD_TABLE_ALIGNMENT.CENTER
title_details.autofit = False
remove_table_borders(title_details)
set_repeat_table_header(title_details.rows[0])
detail_rows = (
    ("Направление", "09.04.02 Информационные системы и технологии"),
    ("Программа", "Интеллектуальные медиатехнологии"),
    ("Номер зачетн. книжки", "2602149"),
    ("Группа", "МИК11"),
    ("Студент", "Рясной Владимир Олегович"),
    ("Преподаватель", "Трубчик Ирина Степановна"),
)
detail_widths = (Inches(2.05), Inches(1.15), Inches(3.30))
for column, width in zip(title_details.columns, detail_widths):
    column.width = width
for row, (label, value) in zip(title_details.rows, detail_rows):
    row.cells[0].text = label
    row.cells[1].text = ""
    row.cells[2].text = value
    for idx, cell in enumerate(row.cells):
        cell.width = detail_widths[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=45, left=0, bottom=45, right=0)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                set_run_font(run, size=14)

add_title_line(doc, "Ростов-на-Дону", size=14, space_before=118, space_after=4)
add_title_line(doc, "2026", size=14)

body_section = doc.add_section(WD_SECTION.NEW_PAGE)
body_section.page_width = Inches(8.27)
body_section.page_height = Inches(11.69)
body_section.top_margin = Inches(0.75)
body_section.bottom_margin = Inches(0.7)
body_section.left_margin = Inches(0.85)
body_section.right_margin = Inches(0.75)
body_section.footer.is_linked_to_previous = False

add_heading(doc, "Краткий результат", 1)
add_body(doc, "Протокол воспроизводимости выполнен полностью: повторные запуски с seed 101 дали одинаковые SHA-256 массивов пикселей, а запуск с seed 1101 — другой отпечаток. При n = 5 увеличение размера мраморной текстуры с 96 до 144 пикселей не обнаружено как существенно влияющее на среднюю яркость, контраст, граничную энергию или энтропию по заранее заданному правилу |Δ| > 2s_base. Предварительная гипотеза не опровергнута.")

add_heading(doc, "1 Цель постановка и гипотеза", 1)
add_body(doc, "Цель работы — подготовить изолированную среду, зафиксировать условия эксперимента в паспорте, подтвердить воспроизводимость по полным выходным данным и оценить влияние одного фактора с учётом разброса между seed.")
add_body(doc, "Вариант 13 использует контекст «мраморный узор для обложки» и фактор «размер изображения x1,5». Исследуется, меняются ли статистики яркости и баланс детальности при увеличении размера с 96 до 144 пикселей при неизменных остальных параметрах.")
add_body(doc, "Гипотеза зафиксирована 24 сентября 2026 года в 16:40:22 MSK до запуска индивидуального варианта. Увеличение размера добавляет случайные узлы решётки, но не изменяет масштаб деталей и закон формирования яркости. Поэтому ожидалось, что для всех четырёх метрик |Δ|/s_base не превысит 2.")
add_body(doc, "Эффект заранее считается существенным, если |Δ| > 2s_base. Это учебное правило, а не формальный статистический тест. Вывод формулируется только для данной модели и n = 5.", bold_lead="Критерий. ")

add_heading(doc, "2 Данные модель и лицензии", 1)
add_body(doc, "Внешние данные не используются. Все изображения синтетически создаются учебным генератором процедурных текстур lab01_generator.py версии 1.0, поэтому отдельная лицензия на входной набор данных не требуется. Случайность локализована в экземпляре random.Random(seed).")
add_body(doc, "Воспроизводимость проверяется по SHA-256 полного массива пикселей. Хеш PNG используется только как дополнительный показатель, поскольку байтовое представление PNG может зависеть от версии библиотеки сжатия zlib.")

add_heading(doc, "3 Среда и параметры", 1)
env = RESULTS["env"]
add_table(doc, ["Параметр", "Значение"], [
    ("Операционная система", env["platform"]),
    ("Архитектура", env["machine"]),
    ("Интерпретатор", f'{env["implementation"]} {env["python"]}'),
    ("zlib", env["zlib"]),
    ("Генератор", f'version {env["generator_version"]}'),
    ("Сторонние пакеты", "не используются; requirements.txt пуст"),
], widths=[Inches(2.0), Inches(4.6)], font_size=9.5)

add_body(doc, "SHA-256 lab01_generator.py: 8f07d340a5a765a8e44f1fec407489cf00114ba54216aa7539c4e7116680e152. SHA-256 lab01_experiment.py: 69257206dab50ef35de448b45764daad0d4f60dbc98febd8953c6faee28e06b5.")

add_table(doc, ["Параметр", "База", "Фактор"], [
    ("Размер", "96 x 96", "144 x 144"),
    ("Шаг решётки cell", "24", "24"),
    ("Число октав", "5", "5"),
    ("Затухание persistence", "0.55", "0.55"),
    ("Квантование", "8 бит", "8 бит"),
    ("Интерполяция", "smooth", "smooth"),
    ("Базовый seed и n", "101; n = 5", "101; n = 5"),
], widths=[Inches(2.4), Inches(2.1), Inches(2.1)], font_size=9.5,
   alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])

add_heading(doc, "4 Ход работы", 1)
add_heading(doc, "4.1 Подготовка среды", 2)
for command in (
    "python3 -m venv .venv",
    ".venv/bin/python --version",
    ".venv/bin/python -m pip freeze",
    ".venv/bin/python -c \"import sys; sys.path.insert(0, 'src'); import lab01_generator\"",
):
    add_command(doc, command)
add_body(doc, "Получен Python 3.12.8; список сторонних пакетов пуст; импорт генератора выполнен успешно.")

add_heading(doc, "4.2 Диагностика отсутствующего seed", 2)
add_table(doc, ["Запуск", "seed", "SHA-256 пикселей начало", "Результат"], [
    ("Без seed 1", "1823359504", "17c1c33f24ddfa5f", "отличается"),
    ("Без seed 2", "1872331504", "bc854907f6fc1805", "отличается"),
    ("seed 101 запуск 1", "101", "93fae5913d797284", "совпадает"),
    ("seed 101 запуск 2", "101", "93fae5913d797284", "совпадает"),
], widths=[Inches(1.5), Inches(1.2), Inches(2.5), Inches(1.4)], font_size=8.8,
   alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
               WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
add_body(doc, "Без явного seed генератор использовал системное время и выдал предупреждение; отпечатки различались. При seed 101 полные массивы и PNG совпали. Причина невоспроизводимости диагностирована: начальное состояние генератора не было зафиксировано.")

add_heading(doc, "4.3 Запуск варианта", 2)
add_command(doc, "python src/lab01_experiment.py --variant 13 --seed 101 --n 5 --out artifacts/v13")
add_body(doc, "Скрипт выполнил контрольный протокол seed 101, 101, 1101, 101, а затем две серии по seed 101…105 для базовых и изменённых параметров.")

add_heading(doc, "5 Результаты", 1)
add_heading(doc, "5.1 Протокол воспроизводимости", 2)
runs = RESULTS["runs_repro"]
add_table(doc, ["Запуск", "seed", "SHA-256 массива начало"], [
    (str(i + 1), str(run["seed"]), run["sha256_pixels"][:16]) for i, run in enumerate(runs)
], widths=[Inches(1.1), Inches(1.2), Inches(4.3)], font_size=9.5,
   alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
add_body(doc, "Получены EQ_1_2 = true, EQ_1_4 = true, EQ_1_3 = false и EQ_2_4 = true. Протокол выполнен полностью. Дополнительно EQ_1_2_png = true в среде zlib 1.2.11.")

add_heading(doc, "5.2 Метрики и влияние фактора", 2)
labels = {"mean": "Средняя яркость", "contrast": "Контраст", "edge": "Граничная энергия", "entropy": "Энтропия"}
metric_rows = []
for key in ("mean", "contrast", "edge", "entropy"):
    t = RESULTS["metrics_table"][key]
    metric_rows.append((labels[key], f'{t["mean_base"]:.4f}', f'{t["s_base"]:.4f}', f'{t["mean_pert"]:.4f}', f'{t["delta"]:+.4f}', f'{t["delta_over_s"]:.2f}', "нет"))
add_table(doc, ["Метрика", "База", "s", "Фактор", "Δ", "|Δ|/s", "Сущ."], metric_rows,
          widths=[Inches(1.65), Inches(0.8), Inches(0.7), Inches(0.8), Inches(0.75), Inches(0.75), Inches(0.65)], font_size=8.2,
          alignments=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6)

for text in (
    "Средняя яркость: отношение 0.71; изменение не обнаружено при n = 5.",
    "Контраст: отношение 0.18; увеличение размера не обнаружено как нарушающее контролируемый контраст фона обложки.",
    "Граничная энергия: отношение 0.26; изменение масштаба локальных переходов не обнаружено.",
    "Энтропия: отношение 0.48; изменение богатства уровней яркости не обнаружено по исходному критерию.",
):
    add_bullet(doc, text)

add_heading(doc, "5.3 Изображения seed 101", 2)
table = doc.add_table(rows=2, cols=2)
table.rows[0].cells[0].text = "База 96 x 96"
table.rows[0].cells[1].text = "Фактор 144 x 144"
for cell, path, alt_text in (
    (table.rows[1].cells[0], BASE_IMAGE, "Базовая мраморная текстура 96 на 96 пикселей с seed 101"),
    (table.rows[1].cells[1], FACTOR_IMAGE, "Мраморная текстура 144 на 144 пикселя после увеличения размера с seed 101"),
):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    picture = run.add_picture(str(path), width=Inches(2.2))
    picture._inline.docPr.set("descr", alt_text)
    picture._inline.docPr.set("title", alt_text)
style_table(table, widths=[Inches(3.3), Inches(3.3)], font_size=9.5,
            alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
p = doc.add_paragraph("Рисунок 1 — Базовая и увеличенная текстуры seed 101 в одинаковом физическом размере")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in p.runs:
    set_run_font(r, size=9, italic=True)

add_heading(doc, "6 Дополнительная проверка при n 50", 1)
rows50 = []
for key in ("mean", "contrast", "edge", "entropy"):
    t = RESULTS_50["metrics_table"][key]
    se = (t["s_base"] ** 2 / 50 + t["s_pert"] ** 2 / 50) ** 0.5
    rows50.append((labels[key], f'{t["delta"]:+.5f}', f'{t["delta_over_s"]:.3f}', f'{se:.5f}', f'{abs(t["delta"])/se:.2f}'))
add_table(doc, ["Метрика", "Δ", "|Δ|/s", "SE разности", "|Δ|/SE"], rows50,
          widths=[Inches(2.0), Inches(1.0), Inches(1.0), Inches(1.3), Inches(1.1)], font_size=8.8,
          alignments=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 4)
add_body(doc, "По заранее заданному правилу 2s_base ни одна метрика не признана существенно изменившейся и при n = 50. Однако стандартная ошибка среднего уменьшается с ростом n. Для энтропии |Δ|/SE ≈ 2.61. Здесь SE — приближение без учёта зависимости между сериями с одинаковыми seed. Для формального вывода нужно учитывать парный дизайн. Эта дополнительная оценка не заменяет и не меняет критерий, зафиксированный до обязательного эксперимента.")

add_heading(doc, "7 Рабочая область Engee", 1)
add_table(doc, ["Параметр", "Наблюдение"], [
    ("Дата и время", "01.10.2026, начало около 10:14 МСК (UTC+3)"),
    ("Версия кабинета", "26.9.1.1; кампусная лицензия ДГТУ"),
    ("Версия среды", "Engee 26.9.2-H1"),
    ("Julia", 'VERSION → v"1.12.4"; versioninfo() выполнена'),
    ("Каталог", "/user/AI_creative_technologies/LR01/"),
    ("Результат", "скрипт сохранён; seed, очистка и перезапуск проверены"),
], widths=[Inches(2.0), Inches(4.6)], font_size=9.2)
add_body(doc, "Часть Б выполнена 1 октября после восстановления доступа. Созданы папки курса и LR01 через контекстное меню «Файлы» → «Создать» → «Папку». Через «+» в «Редакторе скриптов» создан и сохранён LR01_environment.ngscript. В одной ячейке выполнены VERSION и versioninfo(); VERSION дополнительно выведена отдельно. Полный вывод versioninfo() приведён в приложении 12 и сохранён без изменений в artifacts/engee/versioninfo.log.")
add_body(doc, "В отдельной ячейке выполнено seed = 101. В панели «Переменные» появилась строка seed со значением 101, классом Int64 и размером (). После «Очистить все переменные» → «Да» строка исчезла. При этом isdefined(Main, :seed) вернула true, а чтение seed — EngeeNothing(): привязка имени сохранилась, значение было очищено.")
add_body(doc, "После «Перезапуск ядра» → «Перезапустить» таблица очистилась; isdefined(Main, :seed) вернула false. Проверка в командной строке создала служебную переменную ans = false (Bool). Повторное выполнение ячейки seed = 101 в новом ядре восстановило seed. Один повтор после очистки ожидал подключения и был остановлен перед перезапуском; после перезапуска повтор выполнен успешно.")
add_body(doc, "В LR01 через кнопку загрузки панели «Файлы» помещены паспорт части Б и журналы versioninfo.log и state_checks.log. Скрипт сохранён и скачан с результатами. Порядок ячеек: проверка среды, отдельная VERSION, seed. Сохранённые результаты ячеек не равнозначны текущему состоянию памяти; после очистки или перезапуска код объявления переменных нужно выполнить заново.")
add_body(doc, "В кабинете до запуска кнопка называется «Старт Engee», после запуска — «Открыть Engee». Версия кабинета отличается от указанной в методичке. Метаданные экспортированного скрипта содержат шаблонную версию Julia 1.9.3, но выполненные команды показывают 1.12.4. В паспорт внесена фактическая версия выполнения. Ошибки HTTP 502/400/418 от 24 сентября сохранены в журнале как история первой попытки. Учётные данные и персональный URL среды не публикуются.")

add_heading(doc, "8 Ошибки риски и ограничения", 1)
for text in (
    "n = 5 недостаточно для точной оценки разброса и слабых эффектов.",
    "Учебное правило 2s_base не учитывает размер выборки и не является статистическим тестом.",
    "При увеличении изображения генератор создаёт дополнительные узлы решётки; это не масштабирование одного изображения, поэтому корректно сравнивать метрики, а не пиксели.",
    "Исследована одна модель-заменитель и четыре статистики яркости; эстетика и читаемость реального макета непосредственно не измерялись.",
    "Отсутствие обнаруженного эффекта при n = 5 не доказывает полного отсутствия влияния.",
    "Проверка Engee относится к версии 26.9.2-H1 и состоянию 1 октября. Очистка переменных и перезапуск ядра имеют разные последствия; для повторения необходимо соблюдать порядок запуска ячеек.",
):
    add_bullet(doc, text)

add_heading(doc, "9 Ответы на контрольные вопросы", 1)
add_heading(doc, "9.1 Почему одного запуска недостаточно", 2)
add_body(doc, "Один запуск показывает только одну реализацию случайного процесса и не характеризует разброс. Вместо одиночного значения нужно сообщать протокол seed, результаты серии, среднее и выборочное стандартное отклонение. Тогда видно, является ли наблюдаемая разница эффектом фактора или обычным изменением между seed.")
add_heading(doc, "9.2 Почему PNG может иметь другой хеш при тех же пикселях", 2)
add_body(doc, "PNG содержит не только значения пикселей, но и результат сжатия и служебную структуру файла. Разные версии zlib или параметры кодирования могут породить разные байты при одном массиве пикселей. Поэтому основное доказательство воспроизводимости — SHA-256 полного массива пикселей; в паспорте дополнительно записываются версия zlib и хеш PNG.")

add_heading(doc, "10 Вывод", 1)
add_body(doc, "Изолированная среда подготовлена, версии и контрольные суммы кода зафиксированы. Типовая причина невоспроизводимости подтверждена: при пропущенном seed два запуска дают разные полные выходные данные, а при seed 101 — одинаковые. Протокол варианта 13 выполнен полностью. Часть Б Engee также выполнена: подготовлены папки и скрипт, зафиксированы версии и проверено состояние переменных при очистке и перезапуске ядра.")
add_body(doc, "При n = 5 увеличение размера мраморной текстуры с 96 до 144 пикселей не обнаружено как существенно влияющее на среднюю яркость, контраст, граничную энергию и энтропию по критерию |Δ| > 2s_base. Для применения как фона обложки баланс детальности и контролируемый контраст по исследованным метрикам сохранились в пределах разброса между seed. Предварительная гипотеза не опровергнута.")

add_heading(doc, "11 Состав приложенных артефактов", 1)
for text in (
    "artifacts/v13/results.json — полные результаты обязательного эксперимента;",
    "artifacts/v13/passport.json и passport.md — паспорт эксперимента;",
    "artifacts/v13/artifacts — шесть PNG обязательного протокола;",
    "artifacts/error_demo_v13 — четыре запуска диагностики seed;",
    "artifacts/v13_n50 — дополнительная серия n = 50;",
    "experiment_journal_v13.md — журнал с предварительной регистрацией и командами;",
    "reports/LR01_environment.ngscript — экспорт выполненного скрипта Engee с результатами;",
    "artifacts/engee — дополнение к паспорту, полный вывод версии и проверки состояния;",
    "requirements.txt — пустой список сторонних зависимостей.",
):
    add_bullet(doc, text)

heading = add_heading(doc, "12 Полный вывод проверки среды Engee", 1)
heading.paragraph_format.page_break_before = True
for label, value in (("VERSION", 'v"1.12.4"'),
                     ("versioninfo()", (ROOT / "artifacts/engee/versioninfo.log").read_text(encoding="utf-8"))):
    doc.add_paragraph(label)
    p = doc.add_paragraph(value.rstrip("\n"))
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Pt(0)
    for run in p.runs:
        set_run_font(run, name="Courier New", size=8.5)
add_body(doc, "Вывод скопирован из сохранённых результатов скрипта без изменений. Сведения о 112 виртуальных ядрах относятся к тому, что видит Julia на сервере, и не означают выделение всей этой вычислительной мощности пользователю.")

add_page_number(body_section.footer.paragraphs[0])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
